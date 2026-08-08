import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_digital_patwari_schema():
    wb = openpyxl.Workbook()
    
    font_title = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_bold = Font(name='Calibri', size=11, bold=True)
    font_normal = Font(name='Calibri', size=11)
    
    fill_title = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    fill_header = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    fill_accent = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF')
    )
    
    ws1 = wb.active
    ws1.title = "Schema Overview"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.append(["Digital Patwari - Fard / Jamabandi Relational Schema v3"])
    ws1.merge_cells('A1:D1')
    ws1['A1'].font = font_title
    ws1['A1'].fill = fill_title
    ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 35
    
    headers_meta = ["Level", "Entity Name", "Parent Relationship", "Extraction Description / LLM Guidance"]
    ws1.append([])
    ws1.append(headers_meta)
    ws1.row_dimensions[3].height = 24
    
    for col_idx, header in enumerate(headers_meta, 1):
        cell = ws1.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
    meta_data = [
        ("Level 0", "Document Header", "None", "Captures Village name, Hadbast number, Year/Jamabandi cycle, and LR/Legal reference numbers."),
        ("Level 1", "Khewat (खेवट)", "Document Header", "Primary ownership holding group. Vertically merged cells in source documents map to a single Khewat ID."),
        ("Level 2", "Khatauni (खतौनी)", "Khewat", "Cultivation holding group. One Khewat can contain multiple Khataunis."),
        ("Level 3", "Owner / Shareholder", "Khewat", "List of individual owners and their fractional shares (Column 3/4)."),
        ("Level 4", "Khasra (खसरा)", "Khatauni / Owner", "Specific physical parcel/field number with area, soil type, and encumbrances.")
    ]
    
    for row_idx, row_data in enumerate(meta_data, 4):
        ws1.append(list(row_data))
        ws1.row_dimensions[row_idx].height = 20
        for col_idx in range(1, 5):
            c = ws1.cell(row=row_idx, column=col_idx)
            c.font = font_normal
            c.border = thin_border
            if col_idx == 1:
                c.alignment = Alignment(horizontal='center')
                c.font = font_bold

    ws2 = wb.create_sheet(title="Relational Schema")
    ws2.views.sheetView[0].showGridLines = True
    
    tables = [
        ("1. Khewat Table", ["Khewat_ID", "Owner_Group_Ref", "Total_Share", "Remarks_Col8"]),
        ("2. Khatauni Table", ["Khatauni_ID", "Khewat_ID", "Cultivator_Name", "Possession_Type"]),
        ("3. Owner Table", ["Owner_ID", "Khewat_ID", "Owner_Name_Parentage", "Share_Fraction"]),
        ("4. Khasra Table", ["Khasra_Number", "Khatauni_ID", "Area_Kanal_Marla", "Soil_Classification"])
    ]
    
    curr_row = 1
    for table_title, cols in tables:
        ws2.cell(row=curr_row, column=1, value=table_title).font = font_bold
        ws2.cell(row=curr_row, column=1).fill = fill_accent
        ws2.row_dimensions[curr_row].height = 22
        curr_row += 1
        
        ws2.append(cols)
        ws2.row_dimensions[curr_row].height = 20
        for col_idx in range(1, len(cols) + 1):
            c = ws2.cell(row=curr_row, column=col_idx)
            c.font = font_header
            c.fill = fill_header
            c.alignment = Alignment(horizontal='center')
        curr_row += 2

    ws3 = wb.create_sheet(title="Axiom Gates")
    ws3.views.sheetView[0].showGridLines = True
    
    axiom_headers = ["Gate ID", "Rule Name", "Mathematical / Logical Constraint", "Failure Action"]
    ws3.append(axiom_headers)
    ws3.row_dimensions[1].height = 24
    for col_idx, h in enumerate(axiom_headers, 1):
        c = ws3.cell(row=1, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        
    axioms = [
        ("AG-01", "Khewat Share Summation", "Sum of individual Owner share fractions within a Khewat must equal 1.0 (or total share unit).", "Flag extraction error, request re-scan of Column 4."),
        ("AG-02", "Khasra-Khatauni Linkage", "Every Khasra entry must resolve to an active Khatauni container ID.", "Isolate orphaned parcel records."),
        ("AG-03", "Area Consistency Check", "Sum of Khasra areas under a Khatauni must match the recorded total Khatauni area.", "Flag discrepancy for manual Patwari verification.")
    ]
    
    for row_idx, row_data in enumerate(axioms, 2):
        ws3.append(list(row_data))
        ws3.row_dimensions[row_idx].height = 22
        for col_idx in range(1, 5):
            c = ws3.cell(row=row_idx, column=col_idx)
            c.font = font_normal
            c.border = thin_border

    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    output_path = "Digital_Patwari_Schema_v3_CORRECTED.xlsx"
    wb.save(output_path)
    print(f"Successfully saved to current directory: {output_path}")

if __name__ == "__main__":
    create_digital_patwari_schema()
